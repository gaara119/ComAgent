"""Prepare the seven incremental D2 documents and a 13-document manifest.

The six D1 PDFs and all D1 evidence are preserved. Artist summaries are short,
attributed adaptations of browser-verified Wikipedia facts, not downloaded
copies of blocked Wikipedia pages. The PNG is an actual Edge screenshot of
the attributed local summary page. Run DOCX authoring with the bundled Codex
artifact Python runtime; all application/data work stays on project Python 3.12.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/docs/raw/d2"
AUDIT = ROOT / "eval/artifacts/d2/source"

ARTISTS = [
    {"name": "Iron Maiden", "artist_id": 90, "title": "Iron Maiden 铁娘子乐队公开简介", "slug": "iron-maiden", "url": "https://en.wikipedia.org/wiki/Iron_Maiden", "source_document_id": "wikipedia-iron-maiden-introduction",
     "paragraphs": ["Iron Maiden（铁娘子乐队）是英国重金属乐队，由贝斯手及主要创作者 Steve Harris 于1975年在伦敦东部的 Leyton 组建。", "乐队是英国重金属新浪潮的重要代表。其音乐和歌词涉及历史、文学、战争、神话及科幻等题材。", "同名首张录音室专辑于1980年发行；1982年的 The Number of the Beast 是其早期重要作品。这里介绍的是公开音乐背景，不包含 Chinook 销售排名或当代巡演收入。"]},
    {"name": "U2", "artist_id": 150, "title": "U2 乐队公开简介", "slug": "u2", "url": "https://zh.wikipedia.org/wiki/U2%E6%A8%82%E5%9C%98", "source_document_id": "wikipedia-u2-introduction",
     "paragraphs": ["U2是爱尔兰摇滚乐队，于1976年在都柏林成立。其四位核心成员为 Bono、The Edge、Adam Clayton 和 Larry Mullen Jr.。", "乐队的音乐以摇滚为基础，也探索过电子音乐等风格。人权及社会议题常出现在其创作与公共活动中。", "U2于1980年发行首张录音室专辑 Boy，1983年发行 War。这份背景简介不提供未来演出合同、客户数据或实时销售信息。"]},
    {"name": "Metallica", "artist_id": 50, "title": "Metallica 金属乐队公开简介", "slug": "metallica", "url": "https://en.wikipedia.org/wiki/Metallica", "source_document_id": "wikipedia-metallica-introduction",
     "paragraphs": ["Metallica是美国重金属乐队，1981年在洛杉矶由主唱兼吉他手 James Hetfield 和鼓手 Lars Ulrich 组建。其职业生涯的大部分时期以旧金山地区为活动基地。", "乐队以快速节奏、器乐表现和强烈演奏风格著称，被视为鞭笞金属的重要代表之一。", "Master of Puppets 于1986年发行；乐队于2009年进入摇滚名人堂。这里仅提供公开背景，不以音乐史简介解释 Chinook 快照里的商业因果。"]},
    {"name": "Led Zeppelin", "artist_id": 22, "title": "Led Zeppelin 齐柏林飞船乐队公开简介", "slug": "led-zeppelin", "url": "https://zh.wikipedia.org/wiki/%E9%BD%90%E6%9F%8F%E6%9E%97%E9%A3%9E%E8%88%B9%E4%B9%90%E9%98%9F", "source_document_id": "wikipedia-led-zeppelin-introduction",
     "paragraphs": ["Led Zeppelin（齐柏林飞船）是英国摇滚乐队，1968年在伦敦成立。乐队成员包括吉他手 Jimmy Page、主唱 Robert Plant、贝斯兼键盘手 John Paul Jones 和鼓手 John Bonham。", "其早期音乐扎根于蓝调摇滚与迷幻摇滚，也与硬摇滚和重金属的发展密切相关。", "首张同名专辑于1969年发行，1971年的第四张专辑收录了 Stairway to Heaven。此网页是带出处的简短项目改写，不是维基百科原站页面截图。"]},
]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf8")


def render_artist_html(artist: dict) -> str:
    paragraphs = "\n".join(f"<p>{html.escape(p)}</p>" for p in artist["paragraphs"])
    return f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{html.escape(artist['title'])}</title>
<style>body{{max-width:940px;margin:48px auto;padding:0 48px;font:22px/1.8 "Microsoft YaHei",sans-serif;color:#171717;background:white}}h1{{font-size:34px;line-height:1.5}}h2{{font-size:27px;margin:32px 0 12px}}footer{{font-size:16px;border-top:1px solid #ddd;margin-top:32px;padding-top:18px}}a{{overflow-wrap:anywhere;color:#174a80}}</style></head>
<body><main><article><h1>{html.escape(artist['title'])}</h1><h2>公开背景</h2>{paragraphs}
<h2>资料来源与处理说明</h2><p>由项目根据公开百科条目改写为中文简介。用于研究检索和来源定位，不声称它是原站完整页面。</p>
<p>百科原文作者为对应维基百科条目贡献者；本文进行了精简、中文改写与排版，按 CC BY-SA 4.0 共享。未复制外部照片。</p>
<footer>来源 <a href="{html.escape(artist['url'])}">{html.escape(artist['url'])}</a><br>许可 https://creativecommons.org/licenses/by-sa/4.0/<br>Chinook 对应 artist_id {artist['artist_id']}　标准名称 {html.escape(artist['name'])}</footer>
</article></main></body></html>'''


def prepare_scan() -> None:
    """Reproduce the raster source while verifying/reusing existing bytes."""
    import binascii
    import pymupdf
    sys.path.insert(0, str(ROOT))
    from scripts.prepare_rag_data import RemoteZip, REVISION
    entries_path = ROOT / "data/raw/ohr/zip_entries.json"
    source_url = f"https://hf-mirror.com/datasets/opendatalab/OHR-Bench/resolve/{REVISION}/pdfs.zip"
    if entries_path.is_file():
        entries = json.loads(entries_path.read_text(encoding="utf8"))
    else:
        entries = RemoteZip(source_url).entries()
    item = next(item for item in entries if item["name"] == "manual/DUDE_c9c1d919023acc0fe4e48751df45cb5e.pdf")
    native = (RAW / "source/ohr-endoscope-native.pdf").resolve()
    if not native.is_file():
        RemoteZip(source_url).extract(item, native)
    content = native.read_bytes()
    if len(content) != item["size"] or (binascii.crc32(content) & 0xFFFFFFFF) != item["crc32"]:
        raise ValueError("Existing native scan source fails pinned ZIP size/CRC verification")
    scanned = RAW / "ohr-endoscope-scan.pdf"
    with pymupdf.open(native) as original:
        if scanned.is_file():
            with pymupdf.open(scanned) as existing:
                if len(existing) != len(original) or any(page.get_text().strip() or not page.get_images() for page in existing):
                    raise ValueError("Existing scan must have the original page count, raster images and no text layer")
        else:
            with pymupdf.open() as derived:
                for page in original:
                    output_page = derived.new_page(width=page.rect.width, height=page.rect.height)
                    output_page.insert_image(page.rect, pixmap=page.get_pixmap(matrix=pymupdf.Matrix(2, 2)))
                derived.save(scanned, deflate=True)
        audit_path = AUDIT / "ohr-scan-source.json"
        if not audit_path.is_file():
            write_json(audit_path, {"entry": item["name"], "path": native.relative_to(ROOT).as_posix(),
                                   "sha256": sha(native), "size": len(content), "zip_crc32": hex(item["crc32"]),
                                   "source_url": source_url, "source_revision": REVISION,
                                   "derived_scan": scanned.relative_to(ROOT).as_posix(),
                                   "derived_sha256": sha(scanned), "scan_pages": len(original),
                                   "raster_scale": 2, "raster_dpi": 144})


def prepare_sources() -> None:
    import httpx
    import yaml
    from playwright.sync_api import sync_playwright

    RAW.mkdir(parents=True, exist_ok=True)
    AUDIT.mkdir(parents=True, exist_ok=True)
    handoff = json.loads((ROOT / "data/handoff/A_to_B_chinook.json").read_text(encoding="utf8"))
    top4 = [(x["name"], x["artist_id"]) for x in handoff["top_artists"]["rows"][:4]]
    if top4 != [(x["name"], x["artist_id"]) for x in ARTISTS]:
        raise ValueError("D1 SQL Top4 differs from planned document identities; human decision required")
    attempts = []
    with httpx.Client(timeout=20, follow_redirects=True) as client:
        for artist in ARTISTS:
            try:
                response = client.get(artist["url"])
                attempts.append({"url": artist["url"], "http_status": response.status_code,
                                 "response_size": len(response.content), "response_sha256": hashlib.sha256(response.content).hexdigest()})
            except httpx.HTTPError as exc:
                attempts.append({"url": artist["url"], "status": "failed", "error_type": type(exc).__name__})
    for artist in ARTISTS[:2]:
        (RAW / f"{artist['slug']}.html").write_text(render_artist_html(artist), encoding="utf8")
    screenshot_source = RAW / "source/led-zeppelin-summary.html"
    screenshot_source.parent.mkdir(parents=True, exist_ok=True)
    screenshot_source.write_text(render_artist_html(ARTISTS[3]), encoding="utf8")
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, executable_path=r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
        page = browser.new_page(viewport={"width": 1200, "height": 1500}, device_scale_factor=1)
        page.goto(screenshot_source.resolve().as_uri(), wait_until="load")
        page.screenshot(path=str(RAW / "led-zeppelin.png"), full_page=True)
        observed = page.locator("article").inner_text()
        if not all(p in observed for p in ARTISTS[3]["paragraphs"]):
            raise ValueError("Screenshot source did not render the actual article paragraphs")
        screenshot_dimensions = page.evaluate("({width:document.documentElement.scrollWidth,height:document.documentElement.scrollHeight})")
        browser.close()
    write_json(AUDIT / "artist-source-audit.json", {
        "prepared_at_utc": datetime.now(timezone.utc).isoformat(),
        "top_artist_handoff": "data/handoff/A_to_B_chinook.json", "top_artist_handoff_sha256": sha(ROOT / "data/handoff/A_to_B_chinook.json"),
        "direct_fetch_attempts": attempts,
        "adaptation_reason": "Direct local Wikipedia fetch was blocked (HTTP 403); browsing tools provided readable articles. Original total plan allows attributed public introductions when Wikipedia cannot be accessed.",
        "fact_verification_method": "Browser-opened Wikipedia article text; short project-authored Chinese adaptations, not a complete downloaded article",
        "artists": ARTISTS,
        "screenshot": {"path": "data/docs/raw/d2/led-zeppelin.png", "source_html": "data/docs/raw/d2/source/led-zeppelin-summary.html", "sha256": sha(RAW / "led-zeppelin.png"), "browser": "actual headless Microsoft Edge", "is_original_wikipedia_screenshot": False, "rendered_dimensions": screenshot_dimensions},
        "license": "Attributed summaries are adaptations under CC BY-SA 4.0; no external photographs are included"})
    catalog = yaml.safe_load((ROOT / "data/catalog/chinook.yaml").read_text(encoding="utf8"))
    lines = ["# Chinook 数据字典 中文", "", "项目自编字段说明，来源为本项目实际 PostgreSQL 目录。用于理解数据字段，不包含评测问题、标准答案或客户明细。", "", "## 通用约定", "", "数据库为 Chinook 1.4.5，表名和列名采用 snake_case。invoice.invoice_date 为发票日期；已核验快照覆盖2021年至2025年，精确范围见 A/B 数据交接。金额使用数据库金额单位，不推断币种或成功付款状态。", "", "## 地区字段含义", "", "customer.country 表示客户所在地国家。invoice.billing_country 表示发票账单地址国家。两列在当前快照中相等，不能据此将两种业务含义混为同一口径。", "", "## 人名字段含义", "", "customer.first_name 与 employee.first_name 表示名；last_name 表示姓。姓名应显式拼接名与姓。含糊的“名字”不能悄悄改变已确认的题目语义。", "", "## 表与字段", ""]
    for table in catalog["tables"]:
        lines += [f"### {table['name']}", "", table.get("description", ""), "", "| 字段 | 数据类型 | 可为空 | 中文说明 |", "|---|---|---|---|"]
        for column in table["columns"]:
            lines.append("| " + " | ".join([column["name"], column["type"], "是" if column["nullable"] else "否", str(column.get("description", "")).replace("|", "／")]) + " |")
        lines += ["", "主键：" + "、".join(table.get("primary_key", [])) + "。", ""]
    lines += ["## 信息边界", "", "本字典没有客户音乐偏好、银行卡号、实时库存或支付完成状态。字段定义不能替代具体客户事实，也不能提供数据库中不存在的外部指标。", "", "## SQL 注释示例", "", "```sql", "-- # 这一行是代码块里的注释，不是正文标题", "SELECT country FROM public.customer LIMIT 3;", "```", ""]
    (RAW / "chinook-dictionary-zh.md").write_text("\n".join(lines), encoding="utf8")
    write_json(AUDIT / "dictionary-source-audit.json", {"source": "data/catalog/chinook.yaml", "source_sha256": sha(ROOT / "data/catalog/chinook.yaml"), "output": "data/docs/raw/d2/chinook-dictionary-zh.md", "output_sha256": sha(RAW / "chinook-dictionary-zh.md"), "tables": len(catalog["tables"]), "columns": sum(len(t["columns"]) for t in catalog["tables"]), "contains_qa_annotations": False})


def document_style(doc):
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.72)
    section.left_margin = section.right_margin = Inches(.78)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if name == "Normal" else "黑体")
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.12
    # The bundled default OOXML template has a decorative Title border. Remove
    # it explicitly so the native Word render obeys plain black title styling.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(12)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    run = footer.add_run("项目自编研究资料  第 ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页").font.size = Pt(9)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(headers): table.rows[0].cells[i].text = title
    for row in rows:
        for cell, text in zip(table.add_row().cells, row): cell.text = text
    for row_index, row in enumerate(table.rows):
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                element = OxmlElement("w:" + side); element.set(qn("w:w"), "110"); element.set(qn("w:type"), "dxa"); margins.append(element)
            tcpr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                border = OxmlElement("w:" + side); border.set(qn("w:val"), "single"); border.set(qn("w:sz"), "4"); border.set(qn("w:color"), "D9D9D9"); borders.append(border)
            tcpr.append(borders)
            if row_index == 0:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E9EEF4"); tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs: run.font.size = Pt(10.5); run.bold = row_index == 0


def prepare_docx() -> None:
    from docx import Document
    from docx.shared import Pt
    proposal = json.loads((ROOT / "data/rules/metric_definitions.json").read_text(encoding="utf8"))
    artist = ARTISTS[2]
    doc = Document(); document_style(doc)
    doc.add_paragraph("Metallica 金属乐队公开简介", "Title")
    doc.add_paragraph("项目根据公开百科改写的中文研究资料", "Subtitle")
    doc.add_paragraph("本文为 Chinook 艺术家资料准备阶段提供公开音乐背景。对象对应真实 artist_id 50；公开音乐史和本项目数据库销售记录分别保存，本文不声称能够解释销售排名的因果。")
    doc.add_paragraph("公开背景", "Heading 1")
    for paragraph in artist["paragraphs"]: doc.add_paragraph(paragraph)
    doc.add_paragraph("基本信息", "Heading 1")
    add_table(doc, ["字段", "公开信息"], [["成立", "1981年 美国洛杉矶"], ["组建者", "James Hetfield 与 Lars Ulrich"], ["风格", "重金属与鞭笞金属"]], [1.35, 5.5])
    doc.add_paragraph("来源与许可", "Heading 1")
    doc.add_paragraph("出处为维基百科 Metallica 条目及该条目贡献者。项目进行了精简、中文改写和 DOCX 排版，按 CC BY-SA 4.0 共享；未包含外部照片。")
    doc.add_paragraph(artist["url"])
    doc.add_paragraph("https://creativecommons.org/licenses/by-sa/4.0/")
    artist_path = RAW / "metallica-introduction.docx"
    doc.save(artist_path)

    doc = Document(); document_style(doc)
    doc.add_paragraph("Chinook 业务指标口径手册", "Title")
    doc.add_paragraph("项目自编规则文档 D2 版本", "Subtitle")
    approved = proposal["review_status"] == "approved"
    doc.add_paragraph("本手册明确销售额、平均每张发票金额、期间内复购客户和热门曲目的计算方式，用于团队复核 SQL 的字段、粒度和边界。" + ("四项口径已由用户代表 A 和 B 共同确认。" if approved else "四项口径当前待 A 和 B 人工确认，尚未作为在线默认启用。"))
    doc.add_paragraph("范围与通用约定", "Heading 1")
    for p in ["数据来自现有 Chinook 1.4.5 PostgreSQL 快照；已核对可用发票日期范围为2021年1月1日至2025年12月22日。本手册不补写缺失期间，不宣称发票代表成功付款。", "期间统一用 invoice.invoice_date 与左闭右开区间 [开始日, 结束日)。SQL 示例仅以2024年说明过滤写法；具体问答应使用用户明确指定的期间。", "客户所在地使用 customer.country；发票账单地址使用 invoice.billing_country。当前快照中的值相等不改变两种地区含义。", "金额仅表示数据库金额单位，不能推断币种。SUM 与 AVG 在无发票记录时保留 NULL；计数无符合项时为0；零分母返回 NULL 并解释不可计算。"]: doc.add_paragraph(p)
    doc.add_paragraph("指标对应关系", "Heading 1")
    add_table(doc, ["指标全名", "统计粒度", "核心表达式"], [["发票销售额", "发票", "SUM(invoice.total)"], ["平均每张发票金额", "发票", "AVG(invoice.total)"], ["期间内复购客户数", "期间内客户", "发票数至少2张的客户数"], ["热门曲目", "期间内曲目", "SUM(invoice_line.quantity)"]], [1.8, 1.3, 3.75])
    doc.add_paragraph("机器可读规则保存在 data/rules/metric_definitions.json。口径审核与每一道评测题的 gold 人工审核分别记录，互不替代。")
    doc.add_page_break()
    doc.add_paragraph("四项指标的定义与限制", "Heading 1")
    entries = [
        ("销售额", "发票层面按 SUM(invoice.total) 汇总；曲目或艺术家层面按 SUM(invoice_line.unit_price * invoice_line.quantity) 汇总，并连接发票日期进行期间过滤。实际快照已只读核对两种全库总额一致。", "汇总发票总额时不能在发票与明细 JOIN 后重复计入同一张发票。无可用发票的 SUM 结果为 NULL；不得悄悄当作已发生的零销售。"),
        ("平均每张发票金额", "本项目将客单价主口径写全为平均每张发票金额，表达式 AVG(invoice.total)，每张发票权重相同。", "平均每位购买客户金额另命名为 SUM(invoice.total) / COUNT(DISTINCT invoice.customer_id)，不与每张发票平均金额混用。无发票或零分母时保留 NULL。"),
        ("期间内复购客户数", "在用户指定期间内，按 customer_id 分组，统计发票数量至少2张的客户；输出单位为客户数。", "同一客户跨期各有一张发票，不自动算作单期复购。发票资料不能证明付款、退款或商品使用状态。没有合格客户时返回0。"),
        ("热门曲目", "曲目最受欢迎沿用总计划的购买数量默认，即 SUM(invoice_line.quantity)。按销售额或按歌单收录数属于明确的备选口径，不能凭当前排名相同而互换。", "附录中的客户谁买得最多保留金额口径；没有查询对象的短语不套同一个全局默认。并列时以 track_id 升序稳定排序。D2登记口径，不提前实现 D3 完整澄清分级或口径切换。"),
    ]
    for title, definition, limit in entries:
        doc.add_paragraph(title, "Heading 2"); doc.add_paragraph(definition); doc.add_paragraph(limit)
    doc.add_page_break()
    doc.add_paragraph("已核对的 SQL 示例", "Heading 1")
    doc.add_paragraph("以下示例由项目只读账号在同一快照执行并保存原始结果。这里只列写法和边界，不列测试题的具体答案。允许源资料核对所需的子查询，在线 D2 单表问数的安全边界仍由现有 SQL 校验器执行。")
    snippets = [
        ("发票销售额", "SELECT SUM(total) AS revenue\nFROM public.invoice\nWHERE invoice_date >= DATE '2024-01-01'\n  AND invoice_date < DATE '2025-01-01';"),
        ("平均每张发票金额", "SELECT ROUND(AVG(total), 2) AS average_invoice_amount\nFROM public.invoice\nWHERE invoice_date >= DATE '2024-01-01'\n  AND invoice_date < DATE '2025-01-01';"),
        ("期间内复购客户数", "SELECT COUNT(*) AS repeat_customer_count\nFROM (SELECT customer_id FROM public.invoice\n WHERE invoice_date >= DATE '2024-01-01'\n   AND invoice_date < DATE '2025-01-01'\n GROUP BY customer_id HAVING COUNT(*) >= 2) AS r;"),
        ("热门曲目按购买数量", "SELECT track_id, SUM(quantity) AS purchased_quantity\nFROM public.invoice_line\nWHERE invoice_id IN (SELECT invoice_id FROM public.invoice\n WHERE invoice_date >= DATE '2024-01-01'\n   AND invoice_date < DATE '2025-01-01')\nGROUP BY track_id\nORDER BY purchased_quantity DESC, track_id ASC LIMIT 5;"),
    ]
    for title, sql in snippets:
        doc.add_paragraph(title, "Heading 2")
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(sql); r.font.name = "Consolas"; r.font.size = Pt(10)
    doc.add_paragraph("SQL 实跑证据：eval/artifacts/d2/source/metric-sql-checks.json。审核依据：notes/d2/业务指标口径提案.md。", style="Normal")
    manual_path = RAW / "Chinook业务指标口径手册.docx"
    doc.save(manual_path)
    write_json(AUDIT / "docx-authoring-audit.json", {"authoring_runtime": sys.executable, "artifact_kind": "DOCX", "manual_path": manual_path.relative_to(ROOT).as_posix(), "manual_sha256": sha(manual_path), "manual_business_review_status": proposal["review_status"], "artist_docx_path": artist_path.relative_to(ROOT).as_posix(), "artist_docx_sha256": sha(artist_path), "actual_page_count": "pending native Word render", "visual_review": "pending"})


def update_manifest() -> None:
    manifest_path = ROOT / "data/docs/manifest.csv"
    with manifest_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle); original_fields = list(reader.fieldnames or []); previous = list(reader)
    d1_rows = [r for r in previous if r["doc_id"].startswith("ohr-") and r["doc_id"] != "ohr-manual-endoscope-scan"]
    if len(d1_rows) != 6:
        raise ValueError(f"Expected six unchanged D1 PDFs, found {len(d1_rows)}")
    fields = original_fields + [f for f in ("source_document_id", "relative_path", "source_url_or_dataset_id", "license_note", "entity_ref", "is_derived", "content_sha256", "ingestion_status") if f not in original_fields]
    rows = []
    for old in d1_rows:
        row = dict(old)
        row.update(source_document_id=old["doc_id"], relative_path=old["path"], source_url_or_dataset_id=old["source_url"], license_note=old["license"], entity_ref="", is_derived="false", content_sha256=old["content_hash"], ingestion_status="prepared")
        rows.append(row)
    scan_source = json.loads((AUDIT / "ohr-scan-source.json").read_text(encoding="utf8"))
    additions = [
        ("ohr-manual-endoscope-scan", "Rigid Endoscope User Manual 扫描派生件", "pdf", "ohr-endoscope-scan.pdf", "manual", "en", "ohr-manual-endoscope-original", scan_source["source_url"], "OHR research/non-commercial subset; upstream rights retained; derived raster test file", "", "true", "data/docs/raw/d2/source/ohr-endoscope-native.pdf"),
        ("artist-iron-maiden", ARTISTS[0]["title"], "html", "iron-maiden.html", "artist", "zh", ARTISTS[0]["source_document_id"], ARTISTS[0]["url"], "CC BY-SA 4.0 attributed Chinese adaptation; Wikipedia contributors; no external images", "chinook:artist:90:Iron Maiden", "true", ARTISTS[0]["url"]),
        ("artist-u2", ARTISTS[1]["title"], "html", "u2.html", "artist", "zh", ARTISTS[1]["source_document_id"], ARTISTS[1]["url"], "CC BY-SA 4.0 attributed Chinese adaptation; Wikipedia contributors; no external images", "chinook:artist:150:U2", "true", ARTISTS[1]["url"]),
        ("artist-metallica", ARTISTS[2]["title"], "docx", "metallica-introduction.docx", "artist", "zh", ARTISTS[2]["source_document_id"], ARTISTS[2]["url"], "CC BY-SA 4.0 attributed Chinese adaptation; Wikipedia contributors; no external images", "chinook:artist:50:Metallica", "true", ARTISTS[2]["url"]),
        ("artist-led-zeppelin", ARTISTS[3]["title"], "png", "led-zeppelin.png", "artist", "zh", ARTISTS[3]["source_document_id"], ARTISTS[3]["url"], "CC BY-SA 4.0 actual Edge screenshot of attributed project summary HTML, not original Wikipedia page", "chinook:artist:22:Led Zeppelin", "true", "data/docs/raw/d2/source/led-zeppelin-summary.html"),
        ("project-metric-manual", "Chinook 业务指标口径手册 项目自编", "docx", "Chinook业务指标口径手册.docx", "project_rules", "zh", "project-metric-manual-d2", "project-authored; notes/d2/业务指标口径提案.md", "Project-authored team research document; business definitions approved by A+B user attestation", "chinook:business-metrics", "false", ""),
        ("project-chinook-dictionary", "Chinook 数据字典 中文 项目自编", "md", "chinook-dictionary-zh.md", "project_dictionary", "zh", "project-chinook-dictionary-d2", "project-authored from data/catalog/chinook.yaml", "Project-authored team research document; source Chinook catalogue", "chinook:database-dictionary", "false", ""),
    ]
    for doc_id, title, fmt, filename, domain, language, source_id, url, license_note, entity, derived, derived_from in additions:
        path = RAW / filename
        if not path.is_file(): raise FileNotFoundError(path)
        relative = path.relative_to(ROOT).as_posix()
        content_hash = sha(path)
        row = {"doc_id": doc_id, "title": title, "format": fmt, "source_url": url, "license": license_note, "content_hash": content_hash, "path": relative, "domain": domain, "language": language, "derived_from": derived_from, "source_entry": scan_source["entry"] if doc_id.startswith("ohr-") else "", "source_revision": scan_source["source_revision"] if doc_id.startswith("ohr-") else "browser-verified-2026-10-10", "page_count": "2" if doc_id.endswith("-scan") else "", "source_document_id": source_id, "relative_path": relative, "source_url_or_dataset_id": url, "license_note": license_note, "entity_ref": entity, "is_derived": derived, "content_sha256": content_hash, "ingestion_status": "prepared"}
        rows.append(row)
    if len(rows) != 13 or len({r["source_document_id"] for r in rows}) != 13:
        raise ValueError("Corpus must contain exactly 13 independent source documents")
    with manifest_path.open("w", encoding="utf8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields); writer.writeheader(); writer.writerows(rows)
    write_json(AUDIT / "corpus-preparation.json", {"documents": len(rows), "independent_source_documents": len({r["source_document_id"] for r in rows}), "formats": {f: sum(r["format"] == f for r in rows) for f in sorted({r["format"] for r in rows})}, "languages": {l: sum(r["language"] == l for r in rows) for l in sorted({r["language"] for r in rows})}, "manifest_sha256": sha(manifest_path), "all_files_exist": True, "all_source_hashes_verified": all(sha(ROOT / r["path"]) == r["content_hash"] for r in rows), "ingestion_claim": "sources prepared; actual parsing/index audit must be run separately", "qa_annotations_added_to_corpus": False})


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--stage", choices=("scan", "sources", "docx", "manifest"), required=True)
    args = parser.parse_args()
    if args.stage == "scan": prepare_scan()
    elif args.stage == "sources": prepare_sources()
    elif args.stage == "docx": prepare_docx()
    else: update_manifest()
    print(json.dumps({"stage": args.stage, "status": "completed"}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
