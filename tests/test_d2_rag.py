"""D2 parser/retrieval/grounding tests; fixtures never count as real evals."""
from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
import json
from pathlib import Path

import pytest

from qa_agent.config import Settings
from qa_agent.rag.answer import CitationError, DocumentAnswer, answer_from_chunks, validate_answer, citation_coverage_diagnostics
from qa_agent.rag.chunker import chunk_pages
from qa_agent.rag.parsers import parse_document, _ocr_text, DocumentError
from qa_agent.trace import TraceRecorder


def registered(path: Path, fmt: str, doc_id="fixture"):
    return {"doc_id": doc_id, "title": "Fixture", "resolved_path": path, "path": str(path),
            "content_hash": sha256(path.read_bytes()).hexdigest(), "format": fmt,
            "source_url": "https://example.invalid/source"}


def evidence(text, evidence_id="E1", doc_id="fixture"):
    return {"evidence_id": evidence_id, "doc_id": doc_id, "doc_title": doc_id,
            "page": None, "section_path": ["facts"], "bbox": None, "text": text}


def answer(text, ids=None):
    return DocumentAnswer(status="SUCCEEDED", answer=text, citation_ids=ids or ["E1"])


def test_docx_preserves_heading_table_and_order(tmp_path):
    from docx import Document
    source = Document()
    source.add_heading("指标", level=1)
    source.add_paragraph("销售额采用发票金额。")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "意义"
    table.cell(1, 0).text = "total"
    table.cell(1, 1).text = "金额"
    source.add_heading("范围", level=1)
    source.add_paragraph("客户国家。")
    path = tmp_path / "manual.docx"
    source.save(path)
    parsed = parse_document(registered(path, "docx"))
    assert len(parsed) == 2
    assert parsed[0]["section_path"] == ["指标"]
    assert "total | 金额" in parsed[0]["text"]
    assert any(b["kind"] == "table" for b in parsed[0]["blocks"])
    assert all(p["page"] is None for p in parsed)
    chunks = chunk_pages(parsed)
    assert chunks[0]["section_path"] == ["指标"] and chunks[-1]["section_path"] == ["范围"]


def test_html_noise_and_tables_have_real_url_and_heading(tmp_path):
    path = tmp_path / "artist.html"
    path.write_text("<html><script>secret_noise</script><nav>menu_noise</nav><main><h1>Band</h1>"
                    "<p>Formed in 1975.</p><table><tr><th>Genre</th><td>Metal</td></tr></table>"
                    "<h2>History</h2><p>Original statement.</p></main></html>", encoding="utf-8")
    parsed = parse_document(registered(path, "html"))
    joined = "\n".join(p["text"] for p in parsed)
    assert "secret_noise" not in joined and "menu_noise" not in joined
    assert "Genre | Metal" in joined and parsed[-1]["section_path"] == ["Band", "History"]
    assert all(p["page"] is None and p["source_url"] for p in parsed)


def test_markdown_fence_hash_does_not_create_heading(tmp_path):
    path = tmp_path / "dictionary.md"
    path.write_text("# 数据字典\n\n正文。\n```python\n# 代码注释\nx = 2024\n```\n\n## 字段\n\n发票。", encoding="utf-8")
    parsed = parse_document(registered(path, "md"))
    assert len(parsed) == 2
    assert parsed[0]["section_path"] == ["数据字典"]
    assert "# 代码注释" in parsed[0]["text"]
    assert parsed[1]["section_path"] == ["数据字典", "字段"]


def test_pdf_native_text_does_not_invoke_ocr(tmp_path, monkeypatch):
    import pymupdf
    import qa_agent.rag.parsers as module
    path = tmp_path / "native.pdf"
    source = pymupdf.open()
    page = source.new_page()
    page.insert_text((40, 40), "This is a real native PDF document with revenue 1234.")
    source.save(path)
    monkeypatch.setattr(module, "_ocr_engine", lambda: pytest.fail("native PDF must not use OCR"))
    parsed = parse_document(registered(path, "pdf"))
    assert parsed[0]["page"] == 1 and "1234" in parsed[0]["text"]
    assert parsed[0]["parser"] == "pdfplumber+PyMuPDF"


def test_ocr_legacy_result_adapter_coordinates_and_warning(monkeypatch):
    import qa_agent.rag.parsers as module
    engine = lambda image: ([[[[2, 4], [10, 4], [10, 8], [2, 8]], "中国 2024", .7]], [1, 1, 1])
    monkeypatch.setattr(module, "_ocr_engine", lambda: engine)
    text, blocks, warnings = _ocr_text("fixture", scale_x=.5, scale_y=.5)
    assert text == "中国 2024" and blocks[0]["bbox"] == [1, 2, 5, 4]
    assert blocks[0]["char_start"] == 0 and blocks[0]["ocr_confidence"] == .7
    assert warnings and "confidence" in warnings[0]


def test_png_ocr_locator_is_not_pdf_page(tmp_path, monkeypatch):
    from PIL import Image
    import qa_agent.rag.parsers as module
    path = tmp_path / "image.png"
    Image.new("RGB", (100, 30), "white").save(path)
    monkeypatch.setattr(module, "_ocr_text", lambda image: ("原始 OCR 文本。", [
        {"block_id": 0, "text": "原始 OCR 文本。", "bbox": [1, 2, 90, 20],
         "char_start": 0, "char_end": 10, "ocr_confidence": .9}], []))
    parsed = parse_document(registered(path, "png"))
    assert parsed[0]["page"] is None and parsed[0]["coordinate_system"].startswith("image pixels")
    assert chunk_pages(parsed)[0]["bbox"] == [1, 2, 90, 20]


def test_empty_markdown_produces_no_fake_chunks(tmp_path):
    path = tmp_path / "empty.md"
    path.write_text("\n\n", encoding="utf-8")
    assert chunk_pages(parse_document(registered(path, "md"))) == []


def test_section_boundaries_prevent_overlap_and_repeated_sections_have_unique_ids():
    pages = [{"doc_id": "d", "doc_title": "D", "page": None, "record_id": f"s{i}",
              "section_path": [f"Section {i}"], "text": "same source passage "*50, "bbox": None}
             for i in range(2)]
    chunks = chunk_pages(pages, 500, 80)
    assert len({c["chunk_id"] for c in chunks}) == len(chunks)
    assert {tuple(c["section_path"]) for c in chunks} == {("Section 0",), ("Section 1",)}
    assert chunks == chunk_pages(pages, 500, 80)


def test_heading_in_middle_of_pdf_does_not_merge_other_section():
    source = "1. Revenue\nRevenue evidence.\n2. Costs\nCost evidence."
    chunks = chunk_pages([{"doc_id": "d", "doc_title": "D", "page": 1, "text": source, "bbox": None}])
    assert len(chunks) == 2
    assert all(not ("Revenue evidence" in c["text"] and "Cost evidence" in c["text"]) for c in chunks)
    assert all(source[c["char_start"]:c["char_end"]] == c["text"] for c in chunks)


@pytest.mark.parametrize("text,source", [
    ("金额1234.00 [E1]", "金额 1,234.00"),
    ("比例１５％ [E1]", "比例 15%"),
    ("利润-25 [E1]", "利润 −25"),
    ("成立于1975年 [E1]", "Formed in 1975."),
    ("最低保修2年 [E1]", "Endoscopes comewitha2 yearminimumwarranty."),
    ("1. 金额1234 [E1]（第7页）", "Revenue 1234"),
])
def test_numeric_display_normalization_preserves_values(text, source):
    assert validate_answer(answer(text), [evidence(source)])


@pytest.mark.parametrize("text,source", [
    ("比例15% [E1]", "比例 0.15"),
    ("比例15 [E1]", "比例 15%"),
    ("金额25 [E1]", "金额 -25"),
    ("金额$25 [E1]", "金额 €25"),
    ("年度2024 [E1]", "年度2023"),
    ("金额1.2 million [E1]", "金额1.2 billion"),
    ("金额-$25 [E1]", "金额$25"),
])
def test_invalid_numeric_fact_is_blocked(text, source):
    with pytest.raises(CitationError) as error:
        validate_answer(answer(text), [evidence(source)])
    assert error.value.code == "rag_number_invalid"


def test_number_in_other_uncited_chunk_cannot_validate_answer():
    with pytest.raises(CitationError):
        validate_answer(answer("金额9876 [E1]"), [evidence("金额1234"), evidence("金额9876", "E2")])


def test_sentence_citations_cannot_authorize_previous_different_sentence():
    with pytest.raises(CitationError):
        validate_answer(answer("金额9876 [E1]。金额1234 [E2]。", ["E1", "E2"]),
                        [evidence("金额1234"), evidence("金额9876", "E2")])


def test_validation_failure_retries_once_with_feedback_and_true_trace():
    class Model:
        def __init__(self):
            self.calls = []
        def complete_json(self, system, user, cls, trace):
            self.calls.append(json.loads(user))
            return answer("1234 [E99]", ["E99"]) if len(self.calls) == 1 else answer("1234 [E1]")
    model, trace = Model(), TraceRecorder()
    result, used = answer_from_chunks("revenue", [evidence("1234")], model, trace)
    assert result.status == "SUCCEEDED" and len(model.calls) == 2
    assert "validation_feedback" in model.calls[1]
    assert [s.name for s in trace.steps].count("rag_retry") == 1
    assert [s.status for s in trace.steps if s.name == "rag_validate"] == ["error", "ok"]


def test_numeric_validation_failure_after_one_retry_is_visibly_unverified():
    class Model:
        calls = 0
        def complete_json(self, *args):
            self.calls += 1
            return answer("9999 [E1]")
    model, trace = Model(), TraceRecorder()
    result, used = answer_from_chunks("revenue", [evidence("1234")], model, trace)
    assert result.status == "SUCCEEDED" and result.answer.startswith("【未核实】")
    assert result._validation_warnings[0]["code"] == "rag_number_invalid"
    assert used == [evidence("1234")]
    assert model.calls == 2 and [s.name for s in trace.steps].count("rag_regenerate") == 1
    assert [s.status for s in trace.steps if s.name == "rag_validate"] == ["error", "error"]
    assert any(s.name == "rag_unverified" and s.status == "warn" for s in trace.steps)


def test_empty_retrieval_calls_no_model():
    class NoCalls:
        def complete_json(self, *args):
            pytest.fail("No evidence must not call model")
    result, used = answer_from_chunks("unknown", [], NoCalls(), TraceRecorder())
    assert result.status == "EMPTY_RESULT" and used == []


def test_insufficient_status_does_not_display_uncited_model_memory():
    class Model:
        def complete_json(self, *args):
            return DocumentAnswer(status="EMPTY_RESULT", answer="公司2024年收入9999，当前资料证据不足。")
    result, used = answer_from_chunks("2024收入", [evidence("2023年收入1234")], Model(), TraceRecorder())
    assert result.status == "EMPTY_RESULT" and used == []
    assert "9999" not in result.answer and "无法核实" in result.answer


def test_uncited_non_numeric_statement_is_explicit_warning_not_fake_faithfulness():
    result = answer("乐队来自英国[E1]。它的成员来自另外一个国家但未标引用。")
    diagnostic = citation_coverage_diagnostics(result)
    assert diagnostic["potential_uncited_statement_count"] == 1
    assert diagnostic["semantic_claim_coverage"] == "requires_human_review"
    assert diagnostic["mode"] == "heuristic_warning_only"


def test_rrf_top20_each_rank1_formula_and_stable_tie(tmp_path, monkeypatch):
    import csv
    import numpy as np
    import qa_agent.rag.index as index_module
    import qa_agent.rag.embeddings as embeddings
    from qa_agent.documents import load_documents
    from qa_agent.rag.retriever import retrieve
    docs = tmp_path / "data/docs"
    docs.mkdir(parents=True)
    rows = []
    for doc_id in ("allowed", "forbidden"):
        path = docs / f"{doc_id}.pdf"
        path.write_bytes(b"parser fixture only")
        rows.append({"doc_id": doc_id, "title": "Same title", "format": "pdf",
                     "path": f"data/docs/{doc_id}.pdf", "content_hash": sha256(path.read_bytes()).hexdigest()})
    manifest = docs / "manifest.csv"
    with manifest.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    monkeypatch.setattr(index_module, "load_documents", lambda source: load_documents(source, tmp_path))
    monkeypatch.setattr(index_module, "parse_pdf", lambda document: [
        {"doc_id": document["doc_id"], "doc_title": "Same title", "page": i+1,
         "text": "matching revenue evidence", "bbox": None} for i in range(30)])
    class FixtureEncoder:
        def encode(self, texts):
            return np.asarray([[1., 0.]]*len(texts), dtype=np.float32)
        def metadata(self):
            return {"model_id": "test-only"}
    monkeypatch.setattr(embeddings, "get_embedder", lambda settings: FixtureEncoder())
    settings = Settings(db_url="", rag_manifest_path=manifest, rag_index_path=tmp_path / "data/index",
                        rag_top_k=8, rag_candidate_k=20, rag_rrf_k=60)
    index_module.build_index(settings=settings)
    chunks, metadata = retrieve("revenue evidence", "allowed", settings=settings)
    rankings = metadata["rankings"]
    assert metadata["candidate_count"] == 30 and len(chunks) == 8
    assert len(rankings["bm25"]) == len(rankings["dense"]) == 20
    assert all(r["doc_id"] == "allowed" for v in rankings.values() for r in v)
    assert rankings["bm25"][0]["rank"] == 1
    assert rankings["rrf"][0]["score"] == round(2/61, 8)
    assert rankings["bm25"] == sorted(rankings["bm25"], key=lambda r: r["chunk_id"])
    assert metadata["hit_at_5_scope"] == [c["chunk_id"] for c in chunks[:5]]
    np.save(settings.rag_index_path / "embeddings.npy", np.ones((59, 2), dtype=np.float32), allow_pickle=False)
    with pytest.raises(DocumentError) as error:
        index_module.load_index(settings=settings)
    assert error.value.code == "rag_index_corrupt"

